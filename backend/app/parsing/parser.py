"""解析层：把多种格式的源文件解析为统一的中间表示（IR）。

IR 结构：
    IRRBlock = {"type": "text" | "table", "content": str | list[list[str]], "level": int}

设计（见 plan.md）：
- 解析器不直接产出 chunk，先产出 IR，切分层只消费 IR。
- 表格必须被单独识别，作为 "table" 块，供切分层做"整表一块"的旁路处理。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Literal, TypedDict


class IRBlock(TypedDict):
    type: Literal["text", "table"]
    content: str | list[list[str]]
    level: int  # 标题层级；0 表示正文（非标题），1~6 表示 h1~h6


# ------------------------------------------------------------------ Markdown

_MD_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|?\s*$")


def parse_md(text: str) -> list[IRBlock]:
    blocks: list[IRBlock] = []
    lines = text.split("\n")
    i = 0
    para: list[str] = []

    def flush_para() -> None:
        if para:
            joined = "\n".join(para).strip()
            if joined:
                blocks.append({"type": "text", "content": joined, "level": 0})
            para.clear()

    while i < len(lines):
        line = lines[i]
        # 表格：连续的以 | 开头的行
        if line.lstrip().startswith("|"):
            flush_para()
            table_rows: list[list[str]] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw = lines[i].strip()
                if _MD_TABLE_SEP.match(raw):
                    i += 1
                    continue
                cells = [c.strip() for c in raw.strip("|").split("|")]
                table_rows.append(cells)
                i += 1
            if table_rows:
                blocks.append({"type": "table", "content": table_rows, "level": 0})
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_para()
            blocks.append({"type": "text", "content": m.group(2).strip(), "level": len(m.group(1))})
        elif line.strip() == "":
            flush_para()
        else:
            para.append(line)
        i += 1

    flush_para()
    return blocks


# ------------------------------------------------------------------ HTML

def parse_html(text: str) -> list[IRBlock]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(text, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()

    blocks: list[IRBlock] = []
    body = soup.body or soup

    for el in body.descendants:
        name = getattr(el, "name", None)
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            txt = el.get_text(" ", strip=True)
            if txt:
                blocks.append({"type": "text", "content": txt, "level": int(name[1])})
        elif name == "table":
            rows: list[list[str]] = []
            for tr in el.find_all("tr"):
                cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
                if cells:
                    rows.append(cells)
            if rows:
                blocks.append({"type": "table", "content": rows, "level": 0})
        elif name == "p":
            txt = el.get_text(" ", strip=True)
            if txt:
                blocks.append({"type": "text", "content": txt, "level": 0})

    # 去重（descendants 可能重复捕获嵌套结构）
    return _dedupe(blocks)


# ------------------------------------------------------------------ DOCX

def parse_docx(path: Path) -> list[IRBlock]:
    import docx

    doc = docx.Document(str(path))
    blocks: list[IRBlock] = []

    # 按文档顺序遍历段落与表格
    from docx.oxml.ns import qn

    body = doc.element.body
    para_map = {p._element: p for p in doc.paragraphs}
    table_map = {t._element: t for t in doc.tables}

    for child in body.iterchildren():
        if child.tag == qn("w:p") and child in para_map:
            p = para_map[child]
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name or "").lower()
            level = 0
            if style.startswith("heading"):
                m = re.search(r"(\d+)", style)
                level = int(m.group(1)) if m else 1
            elif style in ("title",):
                level = 1
            blocks.append({"type": "text", "content": txt, "level": level})
        elif child.tag == qn("w:tbl") and child in table_map:
            t = table_map[child]
            rows: list[list[str]] = []
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append({"type": "table", "content": rows, "level": 0})
    return blocks


# ------------------------------------------------------------------ PDF

def parse_pdf(path: Path) -> list[IRBlock]:
    """PDF 解析：用 find_tables 抽表，其余文本按行/段落处理。

    标题层级 PDF 无结构信息，这里用字号启发式粗判（较大字号 -> 标题）。
    """
    import pymupdf

    doc = pymupdf.open(str(path))
    blocks: list[IRBlock] = []
    try:
        for page in doc:
            # 表格区域
            tables = page.find_tables()
            table_bboxes = [t.bbox for t in tables.tables]
            for t in tables.tables:
                rows = [[("" if c is None else str(c).strip()) for c in row] for row in t.extract()]
                rows = [r for r in rows if any(r)]
                if rows:
                    blocks.append({"type": "table", "content": rows, "level": 0})

            # 文本：用 dict 模式拿字号，剔除表格区域内的文本
            d = page.get_text("dict")
            for blk in d.get("blocks", []):
                if blk.get("type") != 0:
                    continue
                bbox = blk.get("bbox", (0, 0, 0, 0))
                if any(_overlap(bbox, tb) for tb in table_bboxes):
                    continue
                line_texts = []
                max_size = 0.0
                for line in blk.get("lines", []):
                    spans = line.get("spans", [])
                    txt = "".join(s.get("text", "") for s in spans).strip()
                    if txt:
                        line_texts.append(txt)
                        max_size = max(max_size, max((s.get("size", 0) for s in spans), default=0))
                joined = " ".join(line_texts).strip()
                if not joined:
                    continue
                level = 0
                if max_size >= 15:
                    level = 1
                elif max_size >= 12.5:
                    level = 2
                blocks.append({"type": "text", "content": joined, "level": level})
    finally:
        doc.close()
    return blocks


def _overlap(a: tuple, b: tuple) -> bool:
    ax0, ay0, ax1, ay1 = a
    bx0, by0, bx1, by1 = b
    return not (ax1 <= bx0 or bx1 <= ax0 or ay1 <= by0 or by1 <= ay0)


# ------------------------------------------------------------------ 分发

def _dedupe(blocks: list[IRBlock]) -> list[IRBlock]:
    out: list[IRBlock] = []
    for b in blocks:
        if out and b == out[-1]:
            continue
        out.append(b)
    return out


def parse_file(path: Path) -> list[IRBlock]:
    """按扩展名分发解析。"""
    suffix = path.suffix.lower()
    if suffix == ".md":
        return parse_md(path.read_text(encoding="utf-8"))
    if suffix in (".html", ".htm"):
        return parse_html(path.read_text(encoding="utf-8"))
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"不支持的格式：{suffix}")
